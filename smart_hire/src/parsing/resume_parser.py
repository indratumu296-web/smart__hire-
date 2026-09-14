from __future__ import annotations

from io import BytesIO
import re


SUPPORTED_SUFFIXES = {"txt", "pdf", "docx"}


def _clean_text(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def parse_txt(payload: bytes) -> str:
    """Decode plain-text resume bytes while tolerating malformed characters."""
    return _clean_text(payload.decode("utf-8-sig", errors="ignore"))


def parse_pdf(payload: bytes) -> str:
    """Extract text from every page of a PDF payload."""
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise ValueError("PDF support needs pypdf. Install the packages in requirements.txt.") from error

    try:
        reader = PdfReader(BytesIO(payload))
        return _clean_text("\n".join(page.extract_text() or "" for page in reader.pages))
    except Exception as error:
        raise ValueError("The uploaded PDF could not be read.") from error


def parse_docx(payload: bytes) -> str:
    """Extract paragraph and table text from a DOCX payload."""
    try:
        from docx import Document
    except ImportError as error:
        raise ValueError("DOCX support needs python-docx. Install the packages in requirements.txt.") from error

    try:
        document = Document(BytesIO(payload))
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        return _clean_text("\n".join(parts))
    except Exception as error:
        raise ValueError("The uploaded DOCX file could not be read.") from error


def extract_resume_text(uploaded_file) -> str:
    filename = str(getattr(uploaded_file, "name", ""))
    if "." not in filename:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT resume.")
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT resume.")
    payload = uploaded_file.getvalue()

    if suffix == "txt":
        return parse_txt(payload)
    if suffix == "pdf":
        return parse_pdf(payload)
    if suffix == "docx":
        return parse_docx(payload)
    raise AssertionError("Supported suffix dispatch is incomplete.")
